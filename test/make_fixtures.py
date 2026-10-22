"""Makes the import test files in test/fixtures (all made-up people). Run once; the outputs are committed.

    python3 test/make_fixtures.py

Needs python-docx and Playwright (for the PDF).
"""
import json
import pathlib
import zipfile

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_TAB_ALIGNMENT

HERE = pathlib.Path(__file__).resolve().parent / "fixtures"
HERE.mkdir(exist_ok=True)


def word_cv():
    d = Document()
    d.add_paragraph("Jordan Ellis", style="Title")
    d.add_paragraph("Registered Nurse", style="Subtitle")
    t = d.add_table(rows=1, cols=3)
    for cell, text in zip(t.rows[0].cells, ["Manchester, UK", "jordan.ellis@example.org", "07700 900789"]):
        cell.text = text
    d.add_heading("Personal Statement", level=1)
    d.add_paragraph("Compassionate registered nurse with six years on busy acute wards, calm under pressure and committed to dignified, patient-centred care.")
    d.add_heading("Employment History", level=1)
    for title, org, dates, bullets in [
        ("Senior Staff Nurse", "Northside General Hospital", "May 2021 – Present", ["Lead a team of 6 nurses on a 28-bed acute medical ward.", "Introduced a falls-prevention checklist that halved falls in a year."]),
        ("Staff Nurse", "Riverside Community Trust", "Sep 2018 – Apr 2021", ["Managed caseloads of up to 12 patients a shift.", "Trained 15 student nurses on placement."]),
    ]:
        p = d.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Pt(450), WD_TAB_ALIGNMENT.RIGHT)
        r = p.add_run(f"{title}, {org}")
        r.bold = True
        p.add_run(f"\t{dates}")
        for b in bullets:
            d.add_paragraph(b, style="List Bullet")
    d.add_heading("Education", level=1)
    p = d.add_paragraph()
    p.add_run("BSc (Hons) Adult Nursing").bold = True
    p.add_run(", University of Salford, 2015 – 2018")
    d.add_paragraph("First Class Honours")
    d.add_heading("Skills", level=1)
    d.add_paragraph("Clinical: IV therapy, wound care, medicines management")
    d.add_paragraph("Systems: EPIC, Cerner, Microsoft Office")
    d.add_heading("Languages", level=1)
    d.add_paragraph("English (native), Polish (fluent)")
    d.save(HERE / "word-cv.docx")


TWO_COLUMN_HTML = """<!doctype html><html><head><meta charset="utf-8"><style>
@page { size: A4; margin: 0 }
body { margin: 0; font-family: Arial, Helvetica, sans-serif; font-size: 10pt; color: #222 }
.wrap { display: grid; grid-template-columns: 62mm 1fr; min-height: 297mm }
aside { background: #eef2f5; padding: 16mm 7mm 16mm 12mm }
main { padding: 16mm 14mm 16mm 9mm }
h1 { font-size: 22pt; margin: 0 0 2mm } h2 { font-size: 10.5pt; text-transform: uppercase; letter-spacing: .06em; margin: 7mm 0 2mm; color: #345 }
.role { font-weight: bold; margin-top: 4mm } .meta { color: #555 } ul { margin: 1mm 0 0 4mm; padding: 0 } p { margin: 0 0 1mm }
</style></head><body><div class="wrap"><aside>
<h1>Alex Morgan</h1><p>UX Designer</p>
<h2>Contact</h2><p>alex.morgan@example.net</p><p>+44 7700 900321</p><p>Glasgow, Scotland</p><p>linkedin.com/in/alexmorgan-example</p>
<h2>Skills</h2><p>Figma</p><p>User research</p><p>Prototyping</p><p>Accessibility</p><p>Design systems</p>
<h2>Languages</h2><p>English (native)</p><p>Gaelic (basic)</p>
</aside><main>
<h2>Profile</h2><p>UX designer who turns messy services into simple ones, with eight years across government and retail. Comfortable leading research and shipping production-ready designs.</p>
<h2>Experience</h2>
<p class="role">Lead UX Designer, Thistle Digital</p><p class="meta">January 2020 – Present · Glasgow</p>
<ul><li>Led research and design for a benefits service used by 2 million people.</li><li>Built a design system adopted by 7 product teams.</li></ul>
<p class="role">UX Designer, Lanark Retail Group</p><p class="meta">March 2016 – December 2019</p>
<ul><li>Redesigned checkout, lifting conversion by 18%.</li><li>Ran weekly usability tests with customers.</li></ul>
<h2>Education</h2><p class="role">MA Interaction Design, Glasgow School of Art</p><p class="meta">2014 – 2015</p>
<p class="role">BA Graphic Design, University of Dundee</p><p class="meta">2010 – 2014</p>
</main></div></body></html>"""


def two_column_pdf():
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        b = p.chromium.launch()
        pg = b.new_page()
        pg.set_content(TWO_COLUMN_HTML)
        pg.pdf(path=str(HERE / "two-column.pdf"), format="A4", print_background=True)
        b.close()


def linkedin_zip():
    files = {
        "Profile.csv": "First Name,Last Name,Maiden Name,Address,Birth Date,Headline,Summary,Industry,Zip Code,Geo Location,Twitter Handles,Websites,Instant Messengers\n"
                       'Sam,Okafor,,,,"Backend Engineer at Brightwater","Engineer who likes boring, reliable systems.",Software,,"Cardiff, Wales",,"[PORTFOLIO:https://samokafor.example.dev]",\n',
        "Positions.csv": "Company Name,Title,Description,Location,Started On,Finished On\n"
                         'Brightwater,Backend Engineer,"Built the payments API in Go.\nCut p99 latency by 40%.","Cardiff, Wales",Mar 2022,\n'
                         'Valleys Tech,Junior Developer,Maintained the booking system.,"Swansea, Wales",Jul 2019,Feb 2022\n',
        "Education.csv": "School Name,Start Date,End Date,Notes,Degree Name,Activities\nCardiff University,2016,2019,,BSc Computer Science,Hackathon society\n",
        "Skills.csv": "Name\nGo\nPostgreSQL\nKubernetes\n",
        "Email Addresses.csv": "Email Address,Confirmed,Primary,Updated On\nsam.okafor@example.com,Yes,Yes,2020-01-01\n",
        "Languages.csv": "Name,Proficiency\nWelsh,Professional working proficiency\n",
    }
    with zipfile.ZipFile(HERE / "linkedin-export.zip", "w", zipfile.ZIP_DEFLATED) as z:
        for name, text in files.items():
            z.writestr(name, text)


def structured():
    fresh = {
        "name": "Robin Hart", "meta": {"format": "FRESH@0.6.0"},
        "info": {"label": "Data Analyst", "brief": "Analyst who explains numbers plainly."},
        "contact": {"email": "robin@example.com", "phone": "01632 960123", "website": "https://robinhart.example"},
        "location": {"city": "York", "country": "UK"},
        "employment": {"history": [{"employer": "Minster Insight", "position": "Data Analyst", "start": "2021-06", "highlights": ["Built weekly sales reporting in Power BI."]}]},
        "education": {"history": [{"institution": "University of Leeds", "area": "Statistics", "studyType": "BSc", "start": "2017", "end": "2020", "grade": "2:1"}]},
        "skills": {"sets": [{"name": "Analysis", "skills": ["SQL", "Python", "Power BI"]}]},
        "languages": [{"language": "French", "level": "Conversational"}],
        "social": [{"network": "GitHub", "user": "robinhart", "url": "https://github.com/robinhart"}],
    }
    (HERE / "fresh.json").write_text(json.dumps(fresh, indent=2))
    reactive = {
        "basics": {"name": "Casey Lin", "headline": "Marketing Manager", "email": "casey@example.com", "phone": "", "location": "Brighton", "url": {"href": "https://caseylin.example", "label": ""}, "picture": {"url": ""}},
        "sections": {
            "summary": {"content": "<p>Marketer with a soft spot for <strong>email</strong>.</p>"},
            "experience": {"items": [{"visible": True, "company": "Seagull Media", "position": "Marketing Manager", "location": "Brighton", "date": "Jan 2020 - Present", "summary": "<ul><li>Grew the newsletter to 40,000 readers.</li><li>Ran a team of three.</li></ul>", "url": {"href": ""}}]},
            "education": {"items": [{"visible": True, "institution": "University of Sussex", "studyType": "BA", "area": "English", "score": "", "date": "2014 - 2017", "summary": ""}]},
            "skills": {"items": [{"visible": True, "name": "Email marketing", "description": "Advanced", "keywords": ["Mailchimp", "Klaviyo"]}]},
            "profiles": {"items": [{"visible": True, "network": "LinkedIn", "username": "caseylin", "url": {"href": "https://linkedin.com/in/caseylin"}}]},
        },
    }
    (HERE / "reactive.json").write_text(json.dumps(reactive, indent=2))


if __name__ == "__main__":
    word_cv()
    linkedin_zip()
    structured()
    two_column_pdf()
    print("fixtures written to", HERE)
